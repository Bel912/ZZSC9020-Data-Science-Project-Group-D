import pandas as pd, numpy as np, os
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, Reference

base='/mnt/data'
# Load
pv=pd.read_csv(base+'/Rooftop_PV_202501_to_202608_Combined.csv')
dm=pd.read_csv(base+'/Operational_Demand_202501_to_202608_Combined.csv')
we=pd.read_csv(base+'/Sydney_Airport_Weather_202508_to_202608_Combined.csv')
sol=pd.concat([pd.read_csv(base+'/Sydney Airport North Daily Solar Exposure 2025.csv'),pd.read_csv(base+'/Sydney Airport North Daily Solar Exposure 2026.csv')],ignore_index=True)
# normalize headers
for d in [pv,dm,we,sol]: d.columns=[c.strip().replace('\\_','_') for c in d.columns]
# filter core
pv=pv[(pv['REGIONID']=='NSW1') & (pv['TYPE'].astype(str).str.upper()=='MEASUREMENT')].copy()
dm=dm[dm['REGIONID']=='NSW1'].copy()
pv['DATETIME']=pd.to_datetime(pv['INTERVAL_DATETIME'],dayfirst=True,errors='coerce')
dm['DATETIME']=pd.to_datetime(dm['INTERVAL_DATETIME'],dayfirst=True,errors='coerce')
pv=pv[['DATETIME','POWER_MW','QUALITY_INDICATOR']].rename(columns={'POWER_MW':'ROOFTOP_PV_MW','QUALITY_INDICATOR':'PV_QUALITY_INDICATOR'})
dm=dm[['DATETIME','OPERATIONAL_DEMAND','OPERATIONAL_DEMAND_ADJUSTMENT','WDR_ESTIMATE']].rename(columns={'OPERATIONAL_DEMAND':'OPERATIONAL_DEMAND_MW'})
# dedupe and merge
pv=pv.sort_values('DATETIME').drop_duplicates('DATETIME',keep='last')
dm=dm.sort_values('DATETIME').drop_duplicates('DATETIME',keep='last')
core=pv.merge(dm,on='DATETIME',how='inner',validate='one_to_one')
core['DATE']=core['DATETIME'].dt.normalize()
# weather
we['DATE']=pd.to_datetime(we['Date'],dayfirst=True,errors='coerce').dt.normalize()
keep_weather=['DATE','Minimum temperature (°C)','Maximum temperature (°C)','Rainfall (mm)','Sunshine (hours)','9am relative humidity (%)','9am cloud amount (oktas)','3pm relative humidity (%)','3pm cloud amount (oktas)']
we=we[[c for c in keep_weather if c in we.columns]].drop_duplicates('DATE')
we=we.rename(columns={'Minimum temperature (°C)':'MIN_TEMP_C','Maximum temperature (°C)':'MAX_TEMP_C','Rainfall (mm)':'RAINFALL_MM','Sunshine (hours)':'SUNSHINE_HOURS','9am relative humidity (%)':'RH_9AM_PCT','9am cloud amount (oktas)':'CLOUD_9AM_OKTAS','3pm relative humidity (%)':'RH_3PM_PCT','3pm cloud amount (oktas)':'CLOUD_3PM_OKTAS'})
sol['DATE']=pd.to_datetime(dict(year=sol['Year'],month=sol['Month'],day=sol['Day']),errors='coerce')
solar_col=[c for c in sol.columns if 'solar exposure' in c.lower()][0]
sol=sol[['DATE',solar_col]].rename(columns={solar_col:'SOLAR_EXPOSURE_MJ_M2'}).drop_duplicates('DATE')
weather=we.merge(sol,on='DATE',how='left')
master=core.merge(weather,on='DATE',how='left')
# features
master['YEAR']=master.DATETIME.dt.year; master['MONTH']=master.DATETIME.dt.month; master['MONTH_NAME']=master.DATETIME.dt.month_name(); master['QUARTER']='Q'+master.DATETIME.dt.quarter.astype(str)
master['DAY_OF_WEEK']=master.DATETIME.dt.day_name(); master['WEEKEND_FLAG']=master.DATETIME.dt.dayofweek.ge(5).astype(int)
master['HOUR']=master.DATETIME.dt.hour; master['MINUTE']=master.DATETIME.dt.minute; master['HALF_HOUR_SLOT']=master.DATETIME.dt.strftime('%H:%M')
master['SEASON']=master['MONTH'].map({12:'Summer',1:'Summer',2:'Summer',3:'Autumn',4:'Autumn',5:'Autumn',6:'Winter',7:'Winter',8:'Winter',9:'Spring',10:'Spring',11:'Spring'})
for lag in [1,2,4]:
 master[f'PV_LAG_{lag}']=master.ROOFTOP_PV_MW.shift(lag); master[f'DEMAND_LAG_{lag}']=master.OPERATIONAL_DEMAND_MW.shift(lag)
master['DELTA_PV_MW']=master.ROOFTOP_PV_MW.diff(); master['DELTA_DEMAND_MW']=master.OPERATIONAL_DEMAND_MW.diff()
master['PV_RAMP_DIRECTION']=np.select([master.DELTA_PV_MW>0,master.DELTA_PV_MW<0],['Increase','Decrease'],'No change')
# overlap only where weather period exists
master=master[(master.DATE>=weather.DATE.min()) & (master.DATE<=weather.DATE.max())].copy()
master.to_csv(base+'/NSW1_Historical_Modelling_Dataset_Week3.csv',index=False,date_format='%Y-%m-%d %H:%M:%S')
# summaries
monthly=master.set_index('DATETIME').resample('MS').agg(Avg_PV_MW=('ROOFTOP_PV_MW','mean'),Avg_Demand_MW=('OPERATIONAL_DEMAND_MW','mean'),Peak_PV_MW=('ROOFTOP_PV_MW','max'),Peak_Demand_MW=('OPERATIONAL_DEMAND_MW','max')).reset_index()
hourly=master.groupby('HALF_HOUR_SLOT',sort=True).agg(Avg_PV_MW=('ROOFTOP_PV_MW','mean'),Avg_Demand_MW=('OPERATIONAL_DEMAND_MW','mean')).reset_index()
seasonal=master.groupby('SEASON').agg(Avg_PV_MW=('ROOFTOP_PV_MW','mean'),Avg_Demand_MW=('OPERATIONAL_DEMAND_MW','mean'),Records=('DATETIME','size')).reset_index()
nums=['ROOFTOP_PV_MW','OPERATIONAL_DEMAND_MW','SOLAR_EXPOSURE_MJ_M2','SUNSHINE_HOURS','MAX_TEMP_C','MIN_TEMP_C','RAINFALL_MM','CLOUD_9AM_OKTAS','CLOUD_3PM_OKTAS','DELTA_PV_MW','DELTA_DEMAND_MW']
corr=master[[c for c in nums if c in master]].corr()
# charts
plt.figure(figsize=(10,5)); plt.plot(monthly.DATETIME,monthly.Avg_PV_MW,label='Average rooftop PV'); plt.ylabel('MW'); plt.xlabel('Month'); plt.title('Monthly Average NSW1 Rooftop PV'); plt.xticks(rotation=45); plt.tight_layout(); plt.savefig(base+'/week3_monthly_pv.png',dpi=160); plt.close()
plt.figure(figsize=(10,5)); plt.plot(monthly.DATETIME,monthly.Avg_Demand_MW,label='Average operational demand'); plt.ylabel('MW'); plt.xlabel('Month'); plt.title('Monthly Average NSW1 Operational Demand'); plt.xticks(rotation=45); plt.tight_layout(); plt.savefig(base+'/week3_monthly_demand.png',dpi=160); plt.close()
plt.figure(figsize=(8,6)); plt.imshow(corr,cmap='coolwarm',vmin=-1,vmax=1); plt.colorbar(label='Correlation'); plt.xticks(range(len(corr)),corr.columns,rotation=90,fontsize=7); plt.yticks(range(len(corr)),corr.index,fontsize=7); plt.title('Week 3 Correlation Matrix'); plt.tight_layout(); plt.savefig(base+'/week3_correlation_matrix.png',dpi=180); plt.close()
# data quality
expected=pd.date_range(master.DATETIME.min(),master.DATETIME.max(),freq='30min'); missing_intervals=len(expected.difference(master.DATETIME))
quality=pd.DataFrame({'Metric':['Modelling start','Modelling end','Rows','Duplicate datetimes','Missing 30-minute intervals','Missing PV','Missing demand','Weather match rate','Solar exposure match rate'], 'Value':[str(master.DATETIME.min()),str(master.DATETIME.max()),len(master),master.DATETIME.duplicated().sum(),missing_intervals,master.ROOFTOP_PV_MW.isna().sum(),master.OPERATIONAL_DEMAND_MW.isna().sum(),f"{master['MAX_TEMP_C'].notna().mean():.1%}",f"{master['SOLAR_EXPOSURE_MJ_M2'].notna().mean():.1%}"]})
quality.to_csv(base+'/Week3_Data_Quality_Summary.csv',index=False)
# dictionary
meta={
'DATETIME':('Half-hour interval ending timestamp','datetime','AEMO PV and demand'), 'ROOFTOP_PV_MW':('NSW1 measured rooftop PV power','MW','AEMO Rooftop PV'), 'PV_QUALITY_INDICATOR':('AEMO quality indicator for PV measurement','ratio','AEMO Rooftop PV'), 'OPERATIONAL_DEMAND_MW':('NSW1 operational demand','MW','AEMO Operational Demand'), 'OPERATIONAL_DEMAND_ADJUSTMENT':('Operational demand adjustment','MW','AEMO Operational Demand'), 'WDR_ESTIMATE':('Wholesale demand response estimate','MW','AEMO Operational Demand'), 'DATE':('Calendar date used for daily weather join','date','Derived'), 'MIN_TEMP_C':('Daily minimum temperature at Sydney Airport','°C','BoM'), 'MAX_TEMP_C':('Daily maximum temperature at Sydney Airport','°C','BoM'), 'RAINFALL_MM':('Daily rainfall at Sydney Airport','mm','BoM'), 'SUNSHINE_HOURS':('Daily sunshine duration','hours','BoM'), 'RH_9AM_PCT':('9am relative humidity','%','BoM'), 'CLOUD_9AM_OKTAS':('9am cloud amount','oktas','BoM'), 'RH_3PM_PCT':('3pm relative humidity','%','BoM'), 'CLOUD_3PM_OKTAS':('3pm cloud amount','oktas','BoM'), 'SOLAR_EXPOSURE_MJ_M2':('Daily global solar exposure at Sydney Airport North','MJ/m²','BoM'), 'YEAR':('Calendar year','year','Derived'), 'MONTH':('Calendar month number','1-12','Derived'), 'MONTH_NAME':('Calendar month name','text','Derived'), 'QUARTER':('Calendar quarter','text','Derived'), 'DAY_OF_WEEK':('Day name','text','Derived'), 'WEEKEND_FLAG':('1 for Saturday/Sunday, otherwise 0','binary','Derived'), 'HOUR':('Hour from interval timestamp','0-23','Derived'), 'MINUTE':('Minute from interval timestamp','0/30','Derived'), 'HALF_HOUR_SLOT':('Half-hour time label','HH:MM','Derived'), 'SEASON':('Australian meteorological season','text','Derived'), 'PV_LAG_1':('PV one interval earlier','MW','Derived'), 'PV_LAG_2':('PV two intervals earlier','MW','Derived'), 'PV_LAG_4':('PV four intervals earlier','MW','Derived'), 'DEMAND_LAG_1':('Demand one interval earlier','MW','Derived'), 'DEMAND_LAG_2':('Demand two intervals earlier','MW','Derived'), 'DEMAND_LAG_4':('Demand four intervals earlier','MW','Derived'), 'DELTA_PV_MW':('Change in PV from prior interval','MW','Derived'), 'DELTA_DEMAND_MW':('Change in demand from prior interval','MW','Derived'), 'PV_RAMP_DIRECTION':('Direction of interval PV change','text','Derived')}
dictionary=pd.DataFrame([[c,*meta.get(c,('','',''))] for c in master.columns],columns=['Variable','Description','Unit / Format','Source'])
dictionary.to_csv(base+'/Week3_Data_Dictionary_v2.csv',index=False)
# Excel package using sampled master data to stay manageable
out=base+'/Week3_EDA_and_Documentation.xlsx'
with pd.ExcelWriter(out,engine='openpyxl') as x:
 quality.to_excel(x,sheet_name='Data Quality',index=False)
 dictionary.to_excel(x,sheet_name='Data Dictionary',index=False)
 monthly.to_excel(x,sheet_name='Monthly Summary',index=False)
 hourly.to_excel(x,sheet_name='Half-Hour Profile',index=False)
 seasonal.to_excel(x,sheet_name='Seasonal Summary',index=False)
 corr.to_excel(x,sheet_name='Correlation Matrix')
 master.head(5000).to_excel(x,sheet_name='Model Data Sample',index=False)
 wb=x.book
 for ws in wb.worksheets:
  ws.sheet_view.showGridLines=False; ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
  for cell in ws[1]: cell.fill=PatternFill('solid',fgColor='17365D'); cell.font=Font(color='FFFFFF',bold=True); cell.alignment=Alignment(wrap_text=True)
  for col in range(1,ws.max_column+1):
   vals=[str(ws.cell(r,col).value or '') for r in range(1,min(ws.max_row,150)+1)]
   ws.column_dimensions[get_column_letter(col)].width=min(max(max(map(len,vals))+2,12),36)
print('MASTER_ROWS',len(master),'START',master.DATETIME.min(),'END',master.DATETIME.max())
print('WEATHER_MATCH',master.MAX_TEMP_C.notna().mean(),'SOLAR_MATCH',master.SOLAR_EXPOSURE_MJ_M2.notna().mean(),'MISSING_INTERVALS',missing_intervals)
